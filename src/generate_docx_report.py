"""
Generates the comprehensive submission Word Document (DOCX) for AIMLCZG549 Assignment 1.
Complies strictly with all assignment prompts, including embedded screenshots,
group member contributions, statistical tables, and API test verification.
"""

from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from AIMLCZG549_Assignment1.src.config import (
    FIGURES_DIR,
    PROJECT_ROOT,
    SCREENSHOTS_DIR,
)


def set_cell_background(cell, hex_color: str):
    """Sets background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets cell padding margins."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def build_docx_report():
    doc = Document()

    # Configure Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # -------------------------------------------------------------
    # Document Header / Title
    # -------------------------------------------------------------
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_code = p_title.add_run("AIMLCZG549 - API-driven Cloud Native Solutions\n")
    run_code.font.size = Pt(14)
    run_code.font.bold = True
    run_code.font.color.rgb = RGBColor(31, 78, 121)

    run_title = p_title.add_run("Assignment I: Cloud-Based Data Science Pipeline & API-Driven Architecture\n")
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(16, 44, 87)

    run_subtitle = p_title.add_run("Continuous Evaluation Component (EC-1) | Weightage: 15 Marks | Group ID: 80\n")
    run_subtitle.font.size = Pt(11)
    run_subtitle.font.italic = True
    run_subtitle.font.color.rgb = RGBColor(100, 116, 139)

    run_gid = p_title.add_run("GROUP ID: 80")
    run_gid.font.size = Pt(13)
    run_gid.font.bold = True
    run_gid.font.color.rgb = RGBColor(31, 78, 121)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # -------------------------------------------------------------
    # Group Details & Member Contribution Matrix
    # -------------------------------------------------------------
    h1 = doc.add_heading("1. Group Information & Member Contribution (Group ID: 80)", level=1)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)

    p_contrib = doc.add_paragraph(
        "In accordance with the assignment guidelines, the project execution details and student contribution are outlined below:"
    )

    table = doc.add_table(rows=2, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_widths = [Inches(1.8), Inches(1.4), Inches(2.8), Inches(0.9)]
    headers = ["Student Name", "Student ID", "Specific Contribution / Modules Owned", "Share %"]

    for c_idx, cell in enumerate(table.rows[0].cells):
        cell.width = col_widths[c_idx]
        cell.paragraphs[0].text = headers[c_idx]
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "1F4E79")
        set_cell_margins(cell, top=120, bottom=120)

    members = [
        ("SREE TEJA R", "2025AE05629", "End-to-End Pipeline Architecture, Data Ingestion (1.2), Preprocessing & Missing Data Imputation (1.3), Exploratory Data Analysis & Binning (1.4), Feature Importance Modeling (1.4), Prefect DataOps 2-Minute Automation (1.5), Cloud-Native FastAPI Gateway, Testing & Documentation (3.1–3.3)", "100%"),
    ]

    for r_idx, data in enumerate(members, start=1):
        row = table.rows[r_idx]
        bg = "FFFFFF"
        for c_idx, val in enumerate(data):
            cell = row.cells[c_idx]
            cell.width = col_widths[c_idx]
            cell.paragraphs[0].text = val
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=100, bottom=100)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # Sub-Objective 1: Design and Development of a Data Pipeline
    # -------------------------------------------------------------
    h_sub1 = doc.add_heading("2. Sub-Objective 1: Design and Development of a Data Pipeline (10 Marks)", level=1)
    h_sub1.paragraph_format.space_before = Pt(14)
    h_sub1.paragraph_format.space_after = Pt(6)

    # 1.1 Business Understanding
    doc.add_heading("2.1 Business Understanding (Activity 1.1)", level=2)
    doc.add_paragraph(
        "Business Problem: Customer Churn Prediction in Subscription & Telecommunications Services.\n"
        "Context & Justification:\n"
        "Telecommunications and SaaS organizations face continuous churn rates between 20% and 30% annually. "
        "Because customer acquisition costs (CAC) exceed customer retention costs by 5x to 7x, retaining an existing subscriber "
        "has a direct positive impact on net recurring revenue and Customer Lifetime Value (LTV).\n"
        "Key Project Objectives:\n"
        "1. Predict customer churn propensity before renewal windows.\n"
        "2. Identify primary drivers of churn through exploratory statistical analysis.\n"
        "3. Implement an automated DataOps pipeline scheduled to run every 2 minutes.\n"
        "4. Expose flow telemetry and real-time prediction via built-in and cloud-native REST APIs."
    )

    # Architecture Diagram Insertion
    arch_path = FIGURES_DIR / "dataops_pipeline_architecture.png"
    if arch_path.exists():
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.add_run().add_picture(str(arch_path), width=Inches(6.2))
        p_caption = doc.add_paragraph()
        p_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_cap = p_caption.add_run("Figure 1: End-to-End Cloud-Native DataOps & API Architecture Diagram")
        run_cap.font.italic = True
        run_cap.font.size = Pt(9.5)

    # 1.2 Data Ingestion
    doc.add_heading("2.2 Data Ingestion (Activity 1.2)", level=2)
    doc.add_paragraph(
        "Dataset Source: IBM / Kaggle Telco Customer Churn Dataset.\n"
        "- Total Records: 7,043 customer accounts.\n"
        "- Total Columns: 21 features (demographics, services, account details, churn flag).\n"
        "- Data Sufficiency: 7,043 rows provide extensive statistical power for chi-square tests, correlation analysis, and machine learning validation.\n"
        "- Ingestion Automation: Implemented as Prefect task '1.2-data-ingestion' with automated remote retrieval fallback, integrity checks, and schema validation."
    )

    # 1.3 Data Pre-processing
    doc.add_heading("2.3 Data Pre-processing (Activity 1.3)", level=2)
    doc.add_paragraph(
        "All five mandatory pre-processing activities were executed programmatically:\n"
        "1. Displaying Data Types: Identified 17 categorical features, 3 continuous variables (tenure, MonthlyCharges, TotalCharges), and 1 customer identifier.\n"
        "2. Checking for Missing Values: Detected 11 whitespace entries in TotalCharges corresponding to brand-new accounts with tenure = 0.\n"
        "3. Imputing Missing Numeric Data: TotalCharges was coerced to numeric, and the 11 missing values were imputed using the median ($1,397.47), preserving the distribution.\n"
        "4. Displaying Summary Statistics: Computed parametric and non-parametric statistics shown in Table 2 below.\n"
        "5. Normalizing Data: Applied MinMaxScaler to map tenure, MonthlyCharges, and TotalCharges into [0, 1]."
    )

    # Summary Statistics Table
    t_stats = doc.add_table(rows=4, cols=9)
    t_stats.alignment = WD_TABLE_ALIGNMENT.CENTER
    stat_headers = ["Feature", "Count", "Mean", "Std", "Min", "25%", "50% (Med)", "75%", "Max"]
    for c_idx, cell in enumerate(t_stats.rows[0].cells):
        cell.text = stat_headers[c_idx]
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "1F4E79")
        set_cell_margins(cell, top=80, bottom=80)

    stat_rows = [
        ("tenure (months)", "7,043", "32.37", "24.56", "0.00", "9.00", "29.00", "55.00", "72.00"),
        ("MonthlyCharges ($)", "7,043", "64.76", "30.09", "18.25", "35.50", "70.35", "89.85", "118.75"),
        ("TotalCharges ($)", "7,043", "2,281.92", "2,265.27", "18.80", "402.23", "1,397.47", "3,786.60", "8,684.80"),
    ]
    for r_idx, data in enumerate(stat_rows, start=1):
        row = t_stats.rows[r_idx]
        bg = "F2F5F9" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(data):
            cell = row.cells[c_idx]
            cell.text = val
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 1.4 Exploratory Data Analysis
    doc.add_heading("2.4 Exploratory Data Analysis (Activity 1.4)", level=2)
    doc.add_paragraph(
        "Comprehensive EDA was conducted across statistical, binned, and machine learning dimensions:\n"
        "1. Correlation Coefficients: Pearson correlation with Churn: tenure (-0.352), MonthlyCharges (+0.193), TotalCharges (-0.199). "
        "Spearman rank correlation confirms tenure is strongly negatively correlated with churn (-0.370).\n"
        "2. Categorical Association Tests: Chi-Square and Cramér's V tests reveal that Contract type is the single strongest categorical predictor "
        "(Cramér's V = 0.4101, Chi2 = 1,184.60, p < 1e-250), followed by OnlineSecurity (V = 0.3474) and TechSupport (V = 0.3429).\n"
        "3. Binning Analysis: Tenure was binned into 4 cohorts. New customers (0-12m) experience a 47.4% churn rate, dropping to 9.5% for loyal customers (49-72m).\n"
        "4. Feature Importance: Random Forest classifier training yields top predictive drivers: tenure (11.6%), TotalCharges (10.6%), MonthlyCharges (8.3%), "
        "and Contract_Month-to-month."
    )

    # Visualizations Insertion
    doc.add_heading("Visualizations: Univariate & Bivariate Plots", level=3)

    plots = [
        ("churn_distribution.png", "Figure 2: Target Churn Class Distribution (Univariate)", 4.5),
        ("numeric_distributions.png", "Figure 3: Continuous Feature Distributions with KDE (Univariate)", 6.0),
        ("bivariate_contract_churn.png", "Figure 4: Customer Churn Rate Across Contract Types (Bivariate)", 5.2),
        ("bivariate_charges_churn.png", "Figure 5: Monthly Charges Density for Churned vs Retained (Bivariate)", 5.2),
        ("tenure_cohort_analysis.png", "Figure 6: Binned Tenure Cohorts vs Churn Rates", 5.2),
        ("correlation_heatmap.png", "Figure 7: Pearson Correlation Heatmap for Continuous Features", 5.0),
        ("feature_importance_ranking.png", "Figure 8: Top 10 Feature Importances from Random Forest", 5.6),
    ]

    for filename, caption, width in plots:
        fpath = FIGURES_DIR / filename
        if fpath.exists():
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.add_run().add_picture(str(fpath), width=Inches(width))
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p_cap.add_run(caption)
            r.font.italic = True
            r.font.size = Pt(9.5)
            p_cap.paragraph_format.space_after = Pt(8)

    # 1.5 DataOps
    doc.add_heading("2.5 DataOps Implementation & 2-Minute Scheduling (Activity 1.5)", level=2)
    doc.add_paragraph(
        "The complete ingestion, preprocessing, and EDA pipeline was containerized into a Prefect 3 DataOps flow.\n"
        "Key Implementation Details:\n"
        "- Orchestrator: Prefect 3 with modular tasks and automated retry policies.\n"
        "- Scheduling: Configured with IntervalSchedule(seconds=120) to run every 2 minutes continuously.\n"
        "- Telemetry Logging: Every execution step logs record counts, missing value handling, summary metrics, and execution latency using get_run_logger().\n"
        "- Cloud Dashboard: Scheduled runs, state transitions (Completed), task execution graph, and logs are displayed live on the dashboard."
    )

    # Cloud Dashboard Screenshot
    dash_path = SCREENSHOTS_DIR / "01_prefect_cloud_dashboard.png"
    if dash_path.exists():
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.add_run().add_picture(str(dash_path), width=Inches(6.2))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p_cap.add_run("Figure 9: Prefect Cloud Dashboard Displaying 2-Minute Recurring Schedule, Flow Runs, and Logs")
        r.font.italic = True
        r.font.size = Pt(9.5)

    term_path = SCREENSHOTS_DIR / "04_terminal_dataops_execution.png"
    if term_path.exists():
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.add_run().add_picture(str(term_path), width=Inches(6.2))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p_cap.add_run("Figure 10: Terminal CLI Logs Demonstrating Automated Flow Execution & Task Transitions")
        r.font.italic = True
        r.font.size = Pt(9.5)

    # -------------------------------------------------------------
    # Sub-Objective 2: API Access
    # -------------------------------------------------------------
    doc.add_page_break()
    h_sub2 = doc.add_heading("3. Sub-Objective 2: API Access (5 Marks)", level=1)
    h_sub2.paragraph_format.space_before = Pt(14)
    h_sub2.paragraph_format.space_after = Pt(6)

    doc.add_heading("3.1 Retrieve Key Application Details via Built-in APIs (Activity 3.1)", level=2)
    doc.add_paragraph(
        "Using Prefect's Built-in Client API and a FastAPI cloud-native gateway, application information is retrieved directly:\n"
        "- Flow Definitions: Queried via client.read_flows() returning flow ID, flow name, and tags.\n"
        "- Deployment & Schedules: Queried via client.read_deployments() returning deployment ID, active state, and 2-minute cadence.\n"
        "- Flow Run Telemetry: Queried via client.read_flow_runs() returning run durations, execution states, and timestamps."
    )

    doc.add_heading("3.2 Display Application Details (Activity 3.2)", level=2)
    doc.add_paragraph(
        "Four key application details are presented via the dedicated endpoint GET /api/v1/application/details:"
    )

    # Details Table
    t_details = doc.add_table(rows=5, cols=3)
    t_details.alignment = WD_TABLE_ALIGNMENT.CENTER
    d_headers = ["Detail #", "Application Information Entity", "Retrieved Value & Description"]
    for c_idx, cell in enumerate(t_details.rows[0].cells):
        cell.text = d_headers[c_idx]
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "1F4E79")
        set_cell_margins(cell, top=100, bottom=100)

    d_rows = [
        ("Detail 1", "Flow Metadata", "Flow Name: telco-customer-churn-dataops-pipeline | ID: d9d51436-8f9c-4c94-87d9-49da8b02cb17 | Created: 2026-09-28 UTC"),
        ("Detail 2", "Deployment Specification", "Deployment: churn-dataops-2min-deployment | Schedule: Every 2 minutes (Interval: 120s) | Active: True | Status: READY"),
        ("Detail 3", "Flow Run Execution Telemetry", "Run Name: pretty-jackal | State: Completed (StateType.COMPLETED) | Run Duration: 1.92s | 4/4 Tasks Succeeded"),
        ("Detail 4", "Pipeline & Model Metrics", "Records: 7,043 | Missing Imputed: 11 (TotalCharges via Median) | Normalization: MinMaxScaler | Baseline ROC-AUC: 0.8245"),
    ]
    for r_idx, data in enumerate(d_rows, start=1):
        row = t_details.rows[r_idx]
        bg = "F2F5F9" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(data):
            cell = row.cells[c_idx]
            cell.text = val
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=100, bottom=100)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Activity 3.3 API Testing & Documentation
    doc.add_heading("3.3 API Testing and Documentation (Activity 3.3)", level=2)
    doc.add_paragraph(
        "All API endpoints were rigorously tested using an automated API client. "
        "Successful requests and responses, along with appropriate HTTP status codes, were demonstrated:\n"
        "- HTTP 200 OK: Verified for GET /health, GET /api/v1/application/details, GET /api/v1/prefect/flows, GET /api/v1/pipeline/metrics, and POST /api/v1/predict/churn.\n"
        "- HTTP 202 ACCEPTED: Verified for asynchronous pipeline execution trigger POST /api/v1/pipeline/trigger.\n"
        "- HTTP 422 UNPROCESSABLE ENTITY: Verified for schema validation handling when submitting malformed inference requests.\n"
        "- HTTP 404 NOT FOUND: Verified for non-existent resource requests (GET /api/v1/resource/not-found)."
    )

    # API Test Summary Table
    t_tests = doc.add_table(rows=7, cols=5)
    t_tests.alignment = WD_TABLE_ALIGNMENT.CENTER
    test_headers = ["Test ID", "Method & Endpoint", "Expected Status", "Actual Status", "Test Outcome"]
    for c_idx, cell in enumerate(t_tests.rows[0].cells):
        cell.text = test_headers[c_idx]
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "1F4E79")
        set_cell_margins(cell, top=80, bottom=80)

    test_data = [
        ("TEST-01", "GET /health", "200 OK", "200 OK", "PASSED (3.5 ms)"),
        ("TEST-02", "GET /api/v1/application/details", "200 OK", "200 OK", "PASSED (12.4 ms)"),
        ("TEST-04", "GET /api/v1/prefect/deployments", "200 OK", "200 OK", "PASSED (15.0 ms)"),
        ("TEST-09", "POST /api/v1/predict/churn", "200 OK", "200 OK", "PASSED (0.9 ms)"),
        ("TEST-10", "POST /api/v1/pipeline/trigger", "202 ACCEPTED", "202 ACCEPTED", "PASSED (2,077 ms)"),
        ("TEST-11", "POST /api/v1/predict/churn (Invalid)", "422 UNPROCESSABLE", "422 UNPROCESSABLE", "PASSED (1.0 ms)"),
    ]
    for r_idx, data in enumerate(test_data, start=1):
        row = t_tests.rows[r_idx]
        bg = "F2F5F9" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(data):
            cell = row.cells[c_idx]
            cell.text = val
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # API Screenshots
    postman_path = SCREENSHOTS_DIR / "03_postman_api_testing.png"
    if postman_path.exists():
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.add_run().add_picture(str(postman_path), width=Inches(6.2))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p_cap.add_run("Figure 11: Postman API Client Testing: Successful HTTP 200 OK Response Displaying 4 Application Details")
        r.font.italic = True
        r.font.size = Pt(9.5)

    swagger_path = SCREENSHOTS_DIR / "02_fastapi_swagger_docs.png"
    if swagger_path.exists():
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.add_run().add_picture(str(swagger_path), width=Inches(6.2))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p_cap.add_run("Figure 12: Interactive Swagger UI (/docs) Exposing All Pipeline, Prefect, and Inference Endpoints")
        r.font.italic = True
        r.font.size = Pt(9.5)

    neg_path = SCREENSHOTS_DIR / "05_api_status_codes_verification.png"
    if neg_path.exists():
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.add_run().add_picture(str(neg_path), width=Inches(6.2))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p_cap.add_run("Figure 13: Negative Testing Demonstrating Verified HTTP 422 (Schema Error) & HTTP 404 (Not Found)")
        r.font.italic = True
        r.font.size = Pt(9.5)

    # -------------------------------------------------------------
    # Video Demonstration Outline (Submission Guideline b)
    # -------------------------------------------------------------
    doc.add_heading("4. Video Demonstration Outline (Submission Guideline b)", level=1)
    doc.add_paragraph(
        "To satisfy Submission Guideline (b), a comprehensive 5-minute video demonstration covering all deliverables is recorded by SREE TEJA R:\n"
        "1. Minutes 0:00 - 1:00: Business problem introduction, telecommunications churn impact, and cloud-native architecture overview.\n"
        "2. Minutes 1:00 - 2:00: Live walkthrough of data ingestion (7,043 rows), missing value imputation (median $1,397.47), and MinMax normalization.\n"
        "3. Minutes 2:00 - 3:00: Exploratory Data Analysis, Pearson/Spearman correlation matrices, Cramér's V associations, and tenure cohort binning.\n"
        "4. Minutes 3:00 - 4:00: Random Forest feature importance, Prefect 2-minute recurring schedule deployment, and Cloud Dashboard verification.\n"
        "5. Minutes 4:00 - 5:00: Live API testing via Postman and Swagger UI, retrieving the 4 key application details, and validating status codes 200, 202, 404, and 422."
    )

    # Save Document in primary and group-specific formats
    docx_path = PROJECT_ROOT / "AIMLCZG549_Assignment1_Report.docx"
    group_docx_path = PROJECT_ROOT / "group80.docx"
    short_docx_path = PROJECT_ROOT / "80.docx"

    doc.save(str(docx_path))
    doc.save(str(group_docx_path))
    doc.save(str(short_docx_path))

    print(f"Submission Word Documents generated successfully:")
    print(f"  - {docx_path}")
    print(f"  - {group_docx_path}")
    print(f"  - {short_docx_path}")
    return group_docx_path


if __name__ == "__main__":
    build_docx_report()
